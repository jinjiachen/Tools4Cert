#!/bin/python
#coding:utf-8
"""
Author: Michael Jin
Date: 2021-06-08
"""


import os
from docx import Document
from docx.shared import Inches
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from word import find_table

def find_table(tables,search_string):#�ؼ��ʲ��Ҷ�Ӧ���
    table_components=[]
    for table in tables:#����ÿһ�����
        if table.cell(0,1).text==search_string:#�Զ���ؼ������������
            table_components.append(table)
    return table_components


###���ڰ�TRF�ļ�������Զ�������
def fill_number(document,method):
#    '''
#    document: Document��ʵ��
#    method:cells��rows���ַ�����cells�ٶȿ죬���ǲ���ֱ�ۣ�rowsֱ�ۣ����������������ٶ���
#    '''

    tables = document.tables #��ȡ�ĵ��б������б�
    table5 = tables[5]#�����ж���annex s����
    table7 = tables[7] #�����жϲ���ANNEX S��ʼ�ı�
    if method=='cells':
        cells = table5._cells #�жϲ���һ�������������ǵ�һ�������Ŵ���
        total_rows=len(table5.rows)
        j=1
        for i in range(3,total_rows*4+1,4):
            cells[i].text = str(j)
            j = j+1

        cells = table7._cells #�жϲ���һ�������������ǵڶ��������Ŵ���
        total_rows=len(table7.rows)
        j=1
        for i in range(3,total_rows*4+1,4):
            cells[i].text = str(j)
            j = j+1
    elif method=='rows':
        ###�жϲ���һ�������������ǵ�һ�������Ŵ���
        j=1
        total_rows=len(table5.rows)
        for row in range(0,total_rows):
            print(f'{row}/{total_rows}')
            cells=table5.row_cells(row)
            cells[3].text=str(j)
            j=j+1
    
        ###�жϲ���һ�������������ǵڶ��������Ŵ���
        j=1
        total_rows=len(table7.rows)
        for row in range(0,total_rows):
            print(f'{row}/{total_rows}')
            cells=table7.row_cells(row)
            cells[3].text=str(j)
            j=j+1

    document.save(TRF_path[:-5]+'_filled.docx')

####################################### main program #############################################
while True:
    try:
        print("This TDS is for IEC 60335-2-40:2018 in conjunction with IEC 60335-1:2010+A1:2013+A2:2016!")
        print("================================================== Program begin ==================================================")
#        job_no=input("Please input the project number:")
        TRF_path=input('Please input the TRF file:')
        document = Document(TRF_path) #Open the template document
        fill_number(document,'cells')

#        document.save(TRF_path[:-5]+'_filled.docx')

        print("==================================================Program END ==================================================") 
        flag=input("Press Enter to continue! Others to EXIT!")
        if flag!="":
            break
        else:
            os.system("cls")

    except:
        print("==================================================Program END ==================================================") 
        print("Error, please contact Michael.jc.jin@intertek.com")
        break


